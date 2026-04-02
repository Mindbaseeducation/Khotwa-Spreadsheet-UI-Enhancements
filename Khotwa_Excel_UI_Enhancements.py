import streamlit as st
import pandas as pd
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from io import BytesIO

st.set_page_config(page_title="Excel UI Enhancer", layout="wide")
st.title("📊 Excel UI Enhancer (Custom Styled)")

uploaded_file = st.file_uploader("Upload Excel file", type=["xlsx"])


def apply_ui_formatting(ws):
    """
    Apply custom UI formatting:
    - Gopher font
    - Center alignment
    - Dark blue header
    """

    max_row = ws.max_row
    max_col = ws.max_column

    # Styles
    header_font = Font(name="Gopher", bold=True, color="FFFFFF")
    data_font = Font(name="Gopher")

    header_fill = PatternFill(
        start_color="002060",
        end_color="002060",
        fill_type="solid"
    )

    center_align = Alignment(horizontal="center", vertical="center")

    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    # Header Styling
    for col in range(1, max_col + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border

    # Data Styling
    for row in range(2, max_row + 1):
        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = data_font
            cell.alignment = center_align
            cell.border = border

    # Auto column width
    for col in range(1, max_col + 1):
        col_letter = get_column_letter(col)
        max_length = 0

        for row in range(1, max_row + 1):
            val = ws.cell(row=row, column=col).value
            if val:
                max_length = max(max_length, len(str(val)))

        ws.column_dimensions[col_letter].width = max_length + 3


if uploaded_file:
    try:
        st.success("✅ File uploaded successfully!")

        original_wb = load_workbook(uploaded_file, data_only=True)

        new_wb = Workbook()
        new_wb.remove(new_wb.active)

        for sheet_name in original_wb.sheetnames:
            original_ws = original_wb[sheet_name]

            data = list(original_ws.values)
            if not data:
                continue

            df = pd.DataFrame(data[1:], columns=data[0])

            new_ws = new_wb.create_sheet(title=sheet_name)

            # Write headers
            for c_idx, col_name in enumerate(df.columns, 1):
                new_ws.cell(row=1, column=c_idx, value=col_name)

            # Write data
            for r_idx, row in enumerate(df.values, 2):
                for c_idx, value in enumerate(row, 1):
                    new_ws.cell(row=r_idx, column=c_idx, value=value)

            # Apply formatting
            apply_ui_formatting(new_ws)

        # Save output
        output = BytesIO()
        new_wb.save(output)

        st.download_button(
            label="📥 Download Formatted Excel",
            data=output.getvalue(),
            file_name="Formatted_Output.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        st.info("✨ Custom UI styling applied to all sheets!")

    except Exception as e:
        st.error(f"❌ Error: {e}")
